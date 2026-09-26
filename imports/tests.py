"""Smart institutional data import (Step A3).

AI-assisted mapping is exercised only through deterministic mock providers:
no test calls an external service."""
import csv
import io
from pathlib import Path
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APIClient

from core.models import Department, Faculty, School
from coursework.academic.testing import AcademicFixture
from coursework.models import Course, FacultySubjectAssignment, ScholarCourseEnrollment
from identity.authz import workspaces_of
from identity.models import AuditEvent, CapabilityAssignment, Person
from scholars.models import Scholar

from . import mapping
from .mapping import AIImportMapper
from .models import ImportBatch
from .types import TYPES

API = "/api/academic/imports/"
FIXTURES = Path(__file__).parent / "fixtures"
SCHOLAR_HEADER = ["PRN No", "Student Name", "Email ID", "Gender", "Category", "Entry Qualification", "Dept",
                  "Date of Admission"]


def csv_bytes(rows):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8")


def xlsx_bytes(rows, *, title=None):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faculty"
    if title:
        ws.append([title])
        ws.append([])
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def scholar_row(prn, name=None, email=None, dept="D1", admitted="2026-07-10", **kw):
    row = [prn, name or f"Scholar {prn}", email or f"{prn.lower()}@dypiu.ac.in", kw.get("gender", "F"),
           kw.get("category", "Full-time"), kw.get("entry", "MTECH"), dept, admitted]
    return row


# --- AI mapper mocks (deterministic; never network) -------------------------------------------------------------

class RecordingMapper(AIImportMapper):
    name = "mock"
    calls = []

    def suggest(self, *, import_type, fields, headers):
        RecordingMapper.calls.append({"import_type": import_type, "fields": fields, "headers": dict(headers)})
        out = {}
        for col, h in headers.items():
            if h == "Nom complet":
                out[col] = "name"
            elif h == "Mot de passe":
                out[col] = "password"      # not a field: must be ignored
            elif h == "Courriel perso":
                out[col] = "prn"           # already mapped by the deterministic pass: must be ignored
        return out


class FailingMapper(AIImportMapper):
    name = "down"

    def suggest(self, **kwargs):
        raise TimeoutError("rate limited")


class ImportFixture(AcademicFixture):
    def client_for(self, person):
        if person.user_id is None:
            person.user = get_user_model().objects.create_user(f"imp{person.pk}")
            person.save(update_fields=["user"])
        c = APIClient()
        c.force_login(person.user)
        return c

    def upload(self, content, name):
        return SimpleUploadedFile(name, content)

    def analyze(self, person, itype, content, name="data.csv", **extra):
        return self.client_for(person).post(API + "analyze/", {"import_type": itype,
                                                               "file": self.upload(content, name), **extra},
                                            format="multipart")

    def mapping_of(self, analysis):
        return {str(c["column"]): c["field"] for c in analysis["columns"] if c["field"]}

    def preview(self, person, itype, content, name="data.csv", mapping_=None, provision="none", **extra):
        if mapping_ is None:
            a = self.analyze(person, itype, content, name, **extra)
            assert a.status_code == 200, a.content
            mapping_ = self.mapping_of(a.json())
        import json
        return self.client_for(person).post(API + "preview/", {
            "import_type": itype, "file": self.upload(content, name), "mapping": json.dumps(mapping_),
            "provision": provision, **extra}, format="multipart")

    def commit(self, person, batch_id, content, name="data.csv", skip=False):
        return self.client_for(person).post(f"{API}{batch_id}/commit/", {
            "file": self.upload(content, name), "skip_blocking": "true" if skip else "false"}, format="multipart")


# --- mapping ---------------------------------------------------------------------------------------------------

class MappingTests(ImportFixture):
    def test_aliases_map_to_canonical_fields(self):
        itype = TYPES["scholars"]
        result = mapping.deterministic(itype, dict(enumerate(SCHOLAR_HEADER)))
        got = {s.header: (s.field, s.status) for s in result.values()}
        self.assertEqual(got["Student Name"], ("name", "MAPPED"))
        self.assertEqual(got["PRN No"], ("prn", "MAPPED"))
        self.assertEqual(got["Dept"], ("department", "MAPPED"))
        self.assertEqual(got["Email ID"], ("email", "MAPPED"))
        self.assertEqual(got["Date of Admission"], ("admission_date", "MAPPED"))
        for h in ("Candidate Name", "Scholar Name"):
            self.assertEqual(mapping.deterministic(itype, {0: h})[0].field, "name")

    def test_near_miss_is_only_suggested_and_duplicates_are_ambiguous(self):
        itype = TYPES["scholars"]
        fuzzy = mapping.deterministic(itype, {0: "Scholar Nmae"})[0]
        self.assertEqual((fuzzy.field, fuzzy.status, fuzzy.source), ("name", "SUGGESTED", "FUZZY"))
        both = mapping.deterministic(itype, {0: "Email", 1: "E-mail"})
        self.assertEqual([s.status for s in both.values()], ["AMBIGUOUS", "AMBIGUOUS"])
        self.assertTrue(all(s.field is None for s in both.values()))

    @override_settings(IMPORT_AI_MAPPER="imports.tests.RecordingMapper")
    def test_ai_suggestions_are_validated_and_receive_headers_only(self):
        RecordingMapper.calls = []
        content = csv_bytes([["PRN", "Nom complet", "Mot de passe", "Courriel perso"],
                             ["SECRET-PRN-1", "Secret Person", "hunter2", "x@y.z"]])
        r = self.analyze(self.phd, "scholars", content)
        self.assertEqual(r.status_code, 200, r.content)
        cols = {c["header"]: c for c in r.json()["columns"]}
        self.assertEqual((cols["Nom complet"]["field"], cols["Nom complet"]["status"], cols["Nom complet"]["source"]),
                         ("name", "SUGGESTED", "AI"))
        self.assertIsNone(cols["Mot de passe"]["field"])          # unknown field ignored
        self.assertIsNone(cols["Courriel perso"]["field"])        # already-mapped field ignored
        self.assertTrue(r.json()["ai"]["used"])
        sent = repr(RecordingMapper.calls)
        for secret in ("SECRET-PRN-1", "Secret Person", "hunter2"):
            self.assertNotIn(secret, sent)                        # cell values never leave the server

    @override_settings(IMPORT_AI_MAPPER="imports.tests.FailingMapper")
    def test_unavailable_ai_falls_back_to_deterministic(self):
        r = self.analyze(self.phd, "scholars", csv_bytes([SCHOLAR_HEADER + ["Remarks"], scholar_row("P-X1") + [""]]))
        self.assertEqual(r.status_code, 200)
        self.assertIn("unavailable", r.json()["ai"]["note"])
        self.assertEqual({c["header"]: c["status"] for c in r.json()["columns"]}["Remarks"], "UNMAPPED")
        self.assertEqual(r.json()["missing_required"], [])

    @override_settings(IMPORT_AI_MAPPER="no.such.Mapper")
    def test_misconfigured_ai_is_ignored(self):
        r = self.analyze(self.phd, "scholars", csv_bytes([SCHOLAR_HEADER, scholar_row("P-X2")]))
        self.assertEqual((r.status_code, r.json()["ai"]["enabled"]), (200, False))

    def test_header_row_detected_below_title_rows(self):
        content = xlsx_bytes([["Full name", "Email", "Designation", "Dept"], ["Dr A", "a@dypiu.ac.in", "Professor",
                                                                                 "D1"]], title="Faculty list 2026")
        r = self.analyze(self.acad, "faculty", content, "faculty.xlsx")
        self.assertEqual((r.status_code, r.json()["header_row"]), (200, 3))


# --- validation and preview --------------------------------------------------------------------------------------

class ScholarImportTests(ImportFixture):
    def setUp(self):
        super().setUp()
        Department.objects.create(school=self.s1, code="PHY1", name="Physics")
        Department.objects.create(school=self.s2, code="PHY2", name="Physics")
        self.content = csv_bytes([
            SCHOLAR_HEADER,
            scholar_row("NEW001"),                                     # 2 valid
            scholar_row("NEW002", dept="Dept One"),                    # 3 valid (department by exact name)
            scholar_row("NEW001"),                                     # 4 duplicate of row 2
            scholar_row("NEW001", name="Someone Else"),                # 5 conflict with row 2
            scholar_row(self.sch1.prn),                                # 6 conflict with an existing scholar
            scholar_row("NEW003", admitted="31/02/2026"),              # 7 invalid date
            scholar_row("NEW004", email="not-an-email"),               # 8 invalid e-mail
            scholar_row("NEW005", category="PhD"),                     # 9 invalid enum value
            ["NEW006", "", "new006@dypiu.ac.in", "F", "FT", "MTECH", "D1", "2026-07-10"],  # 10 missing name
            scholar_row("NEW007", dept="Physics"),                     # 11 ambiguous department
            scholar_row("NEW008", dept="Chemistry"),                   # 12 unknown department
        ])

    def report(self, response):
        return {r["row"]: r for r in response.json()["report"]}

    def test_preview_classifies_every_row_and_writes_nothing(self):
        before = (Scholar.objects.count(), Person.objects.count(), AuditEvent.objects.filter(
            action="scholar.edit_record.create").count())
        r = self.preview(self.phd, "scholars", self.content)
        self.assertEqual(r.status_code, 201, r.content)
        rep = self.report(r)
        self.assertEqual({n: rep[n]["status"] for n in rep}, {
            2: "VALID", 3: "VALID", 4: "DUPLICATE", 5: "CONFLICT", 6: "CONFLICT", 7: "INVALID", 8: "INVALID",
            9: "INVALID", 10: "MISSING", 11: "AMBIGUOUS", 12: "INVALID"})
        self.assertEqual(r.json()["counts"], {"VALID": 2, "DUPLICATE": 1, "INVALID": 4, "AMBIGUOUS": 1,
                                              "MISSING": 1, "CONFLICT": 2})
        self.assertIn("PHY1", rep[11]["reasons"][0])
        self.assertIn("never merged", rep[6]["reasons"][0])
        self.assertEqual((Scholar.objects.count(), Person.objects.count(), AuditEvent.objects.filter(
            action="scholar.edit_record.create").count()), before)
        self.assertEqual(ImportBatch.objects.get().status, "PREVIEWED")

    def test_commit_requires_explicit_skip_then_imports_valid_rows_atomically(self):
        b = self.preview(self.phd, "scholars", self.content).json()
        refused = self.commit(self.phd, b["id"], self.content)
        self.assertEqual(refused.status_code, 400)
        self.assertFalse(Scholar.objects.filter(prn__startswith="NEW").exists())
        done = self.commit(self.phd, b["id"], self.content, skip=True)
        self.assertEqual(done.status_code, 200, done.content)
        self.assertEqual(sorted(Scholar.objects.filter(prn__startswith="NEW").values_list("prn", flat=True)),
                         ["NEW001", "NEW002"])
        batch = ImportBatch.objects.get(pk=b["id"])
        self.assertEqual((batch.status, batch.rows_created, batch.rows_skipped), ("COMMITTED", 2, 9))
        ev = AuditEvent.objects.get(pk=batch.audit_event_id)
        self.assertEqual((ev.action, ev.allowed, ev.actor_id, ev.after["batch"]),
                         ("imports.scholars.commit", True, self.phd.pk, batch.pk))
        with self.assertRaises(PermissionError):
            batch.rows_created = 99
            batch.save()
        with self.assertRaises(PermissionError):
            batch.delete()
        self.assertEqual(self.commit(self.phd, b["id"], self.content, skip=True).status_code, 400)  # once only

    def test_same_file_twice_creates_nothing_new(self):
        content = csv_bytes([SCHOLAR_HEADER, scholar_row("IDEM01"), scholar_row("IDEM02")])
        b = self.preview(self.phd, "scholars", content).json()
        self.assertEqual(self.commit(self.phd, b["id"], content).status_code, 200)
        again = self.preview(self.phd, "scholars", content).json()
        self.assertEqual(again["counts"]["DUPLICATE"], 2)
        self.assertEqual(self.commit(self.phd, again["id"], content).status_code, 400)  # nothing valid
        self.assertEqual(Scholar.objects.filter(prn__startswith="IDEM").count(), 2)

    def test_data_changed_since_preview_is_a_conflict_and_writes_nothing(self):
        content = csv_bytes([SCHOLAR_HEADER, scholar_row("RACE01"), scholar_row("RACE02")])
        b = self.preview(self.phd, "scholars", content).json()
        Scholar.objects.create(prn="RACE02", name="Other", email="other@dypiu.ac.in", gender="M", category="FT",
                               entry_qualification="MTECH", department=self.d1, admission_date="2026-07-01")
        r = self.commit(self.phd, b["id"], content)
        self.assertEqual(r.status_code, 409)
        self.assertFalse(Scholar.objects.filter(prn="RACE01").exists())
        self.assertEqual(ImportBatch.objects.get(pk=b["id"]).status, "PREVIEWED")

    def test_commit_needs_the_same_file_and_the_same_person(self):
        content = csv_bytes([SCHOLAR_HEADER, scholar_row("SAME01")])
        b = self.preview(self.phd, "scholars", content).json()
        other = csv_bytes([SCHOLAR_HEADER, scholar_row("SAME02")])
        self.assertEqual(self.commit(self.phd, b["id"], other).status_code, 400)
        self.assertEqual(self.commit(self.acad, b["id"], content).status_code, 403)
        self.assertFalse(Scholar.objects.filter(prn__startswith="SAME").exists())

    def test_failure_midway_rolls_back_everything(self):
        content = csv_bytes([SCHOLAR_HEADER, scholar_row("ROLL01"), scholar_row("ROLL02"), scholar_row("ROLL03")])
        b = self.preview(self.phd, "scholars", content).json()
        from coursework.academic import institution
        real = institution.create_scholar_record
        calls = {"n": 0}

        def flaky(*a, **kw):
            calls["n"] += 1
            if calls["n"] == 2:
                raise RuntimeError("database went away")
            return real(*a, **kw)

        with mock.patch("coursework.academic.institution.create_scholar_record", side_effect=flaky):
            c = self.client_for(self.phd)
            c.raise_request_exception = False
            r = c.post(f"{API}{b['id']}/commit/", {"file": self.upload(content, "data.csv")}, format="multipart")
        self.assertEqual(r.status_code, 500)
        self.assertFalse(Scholar.objects.filter(prn__startswith="ROLL").exists())
        self.assertEqual(ImportBatch.objects.get(pk=b["id"]).status, "PREVIEWED")

    def test_large_file_within_limits(self):
        rows = [SCHOLAR_HEADER] + [scholar_row(f"BULK{i:05d}") for i in range(1000)]
        content = csv_bytes(rows)
        b = self.preview(self.phd, "scholars", content).json()
        self.assertEqual(b["counts"]["VALID"], 1000)
        self.assertEqual(self.commit(self.phd, b["id"], content).status_code, 200)
        self.assertEqual(Scholar.objects.filter(prn__startswith="BULK").count(), 1000)

    @override_settings(IMPORT_MAX_ROWS=50)
    def test_row_limit(self):
        content = csv_bytes([SCHOLAR_HEADER] + [scholar_row(f"LIM{i:03d}") for i in range(60)])
        r = self.analyze(self.phd, "scholars", content)
        self.assertEqual(r.status_code, 400)
        self.assertIn("limit", r.json()["detail"])

    @override_settings(IMPORT_MAX_BYTES=100)
    def test_size_limit(self):
        r = self.analyze(self.phd, "scholars", csv_bytes([SCHOLAR_HEADER] + [scholar_row("SZ1")] * 5))
        self.assertEqual(r.status_code, 400)


class FileFormatTests(ImportFixture):
    def test_xlsx_import_with_title_rows_and_provisioning(self):
        content = xlsx_bytes([["Faculty Name", "Official Email", "Designation", "Dept"],
                              ["Dr Xlsx One", "xone@dypiu.ac.in", "Professor", "D1"],
                              ["Dr Xlsx Two", "xtwo@dypiu.ac.in", "Assistant Professor", "Dept Two"]],
                             title="DYPIU faculty")
        grants = CapabilityAssignment.objects.count()
        b = self.preview(self.acad, "faculty", content, "faculty.xlsx", provision="login")
        self.assertEqual(b.status_code, 201, b.content)
        self.assertEqual(b.json()["counts"]["VALID"], 2)
        self.assertFalse(Faculty.objects.filter(email="xone@dypiu.ac.in").exists())
        r = self.commit(self.acad, b.json()["id"], content, "faculty.xlsx")
        self.assertEqual(r.status_code, 200, r.content)
        person = Person.objects.get(email="xone@dypiu.ac.in")
        self.assertFalse(person.user.has_usable_password())
        self.assertEqual([w["key"] for w in workspaces_of(person)], ["faculty"])
        self.assertEqual(CapabilityAssignment.objects.count(), grants)  # derived, never granted
        self.assertEqual(r.json()["provisioned"], 2)

    def test_legacy_xls_import(self):
        content = (FIXTURES / "scholars_legacy.xls").read_bytes()
        a = self.analyze(self.phd, "scholars", content, "scholars.xls").json()
        self.assertEqual((a["header_row"], a["missing_required"], a["rows"]), (3, [], 2))
        b = self.preview(self.phd, "scholars", content, "scholars.xls").json()
        self.assertEqual(b["counts"]["VALID"], 2, b["report"])
        self.assertEqual(self.commit(self.phd, b["id"], content, "scholars.xls").status_code, 200)
        s = Scholar.objects.get(prn="XLS002")
        self.assertEqual((s.gender, s.category, str(s.admission_date)), ("M", "FT", "2026-07-10"))

    def test_malformed_and_unsupported_files(self):
        for content, name in ((b"PK\x03\x04 not really a workbook", "x.xlsx"), (b"\xd0\xcf\x11\xe0 junk", "x.xls"),
                              (b"%PDF-1.4", "x.pdf"), (b"", "x.csv"), (csv_bytes([SCHOLAR_HEADER]), "x.csv"),
                              (b"a,b\x00c\n1,2", "x.csv")):
            r = self.analyze(self.phd, "scholars", content, name)
            self.assertEqual(r.status_code, 400, name)
            self.assertNotIn("Traceback", r.content.decode(errors="ignore"))

    def test_zip_bomb_workbook_is_refused(self):
        content = xlsx_bytes([["Faculty Name", "Email"], ["x" * 2000, "a@dypiu.ac.in"]] * 200)
        with override_settings(IMPORT_MAX_EXPANDED_BYTES=50_000):
            r = self.analyze(self.acad, "faculty", content, "big.xlsx")
        self.assertEqual(r.status_code, 400)
        self.assertIn("expands", r.json()["detail"])

    def test_csv_with_semicolons_and_windows_encoding(self):
        content = "PRN;Full name;Email;Gender;Category;Entry qualification;Department;Admission date\r\n" \
                  "WIN01;Zoë Désai;win01@dypiu.ac.in;F;FT;MTECH;D1;10-07-2026\r\n".encode("cp1252")
        b = self.preview(self.phd, "scholars", content).json()
        self.assertEqual(b["counts"]["VALID"], 1, b["report"])
        self.assertEqual(self.commit(self.phd, b["id"], content).status_code, 200)
        self.assertEqual(Scholar.objects.get(prn="WIN01").name, "Zoë Désai")


# --- other import types --------------------------------------------------------------------------------------------

class OtherTypeTests(ImportFixture):
    def run_import(self, person, itype, rows, **kw):
        content = csv_bytes(rows)
        b = self.preview(person, itype, content, **kw)
        self.assertEqual(b.status_code, 201, b.content)
        return b.json(), content

    def test_structure_import(self):
        b, content = self.run_import(self.acad, "structure", [
            ["School code", "School name", "Department code", "Department name"],
            ["SOM", "School of Management", "MBA", "Business Administration"],
            ["SOM", "", "FIN", "Finance"],                      # school created by the row above
            ["S1", "Renamed School", "D9", "Dept Nine"],        # conflicts with the existing school name
            ["S1", "", "D1", "Dept One"],                       # existing department: duplicate
            ["NEWS", "", "D10", "Dept Ten"],                    # new school without a name
        ])
        self.assertEqual([r["status"] for r in b["report"]], ["VALID", "VALID", "CONFLICT", "DUPLICATE", "MISSING"])
        self.assertEqual(self.commit(self.acad, b["id"], content, skip=True).status_code, 200)
        self.assertEqual(School.objects.get(code="SOM").departments.count(), 2)
        self.assertEqual(School.objects.get(code="S1").name, "School One")  # never changed

    def test_course_import(self):
        b, content = self.run_import(self.acad, "courses", [
            ["Course Code", "Course Title", "Credits", "Category", "Dept"],
            ["NC101", "New Elective", "3", "Elective", "D1"],
            ["SIS7002", "Something else", "4", "MANDATORY", ""],
            ["NC102", "Bad credits", "three", "ELECTIVE", ""],
        ])
        self.assertEqual([r["status"] for r in b["report"]], ["VALID", "CONFLICT", "INVALID"])
        self.commit(self.acad, b["id"], content, skip=True)
        self.assertTrue(Course.objects.filter(code="NC101", department=self.d1).exists())

    def test_enrollment_import_applies_enrollment_rules(self):
        b, content = self.run_import(self.phd, "enrollments", [
            ["PRN", "Course code", "Academic year", "Term", "Section"],
            [self.sch1.prn, "SIS7002", "AY", "JUL_DEC", "A"],
            [self.sch2.prn, "SIS7002", "AY", "July–December", "Z"],   # no such section
            [self.sch2.prn, "SIS7002", "1999-00", "JUL_DEC", "A"],   # no such year
            ["NOPE", "SIS7002", "AY", "JUL_DEC", "A"],                # unknown scholar
        ])
        self.assertEqual([r["status"] for r in b["report"]], ["VALID", "INVALID", "INVALID", "INVALID"])
        self.assertEqual(self.commit(self.phd, b["id"], content, skip=True).status_code, 200)
        self.assertTrue(ScholarCourseEnrollment.objects.filter(scholar=self.sch1, offering=self.offering).exists())
        again, _ = self.run_import(self.phd, "enrollments", [["PRN", "Course code", "Academic year", "Term", "Section"],
                                                             [self.sch1.prn, "SIS7002", "AY", "JUL_DEC", "A"]])
        self.assertEqual(again["report"][0]["status"], "DUPLICATE")

    def test_assignment_import(self):
        b, content = self.run_import(self.acad, "assignments", [
            ["Faculty e-mail", "Course code", "Academic year", "Term", "Role", "Section", "Approval basis"],
            [self.outsider.email, "SIS7002", "AY", "JUL_DEC", "Instructor", "A", "Order 1"],
            [self.outsider.email, "SIS7002", "AY", "JUL_DEC", "Coordinator", "", "Order 2"],  # already has one
        ])
        # Faculty records have no e-mail in the fixture: use the person e-mail on the record for this test.
        self.assertEqual(b["report"][0]["status"], "INVALID")  # no faculty record with that e-mail
        fac = self.outsider.faculty_profile
        fac.email = self.outsider.email
        fac.save()
        b, content = self.run_import(self.acad, "assignments", [
            ["Faculty e-mail", "Course code", "Academic year", "Term", "Role", "Section", "Approval basis"],
            [self.outsider.email, "SIS7002", "AY", "JUL_DEC", "Instructor", "A", "Order 1"],
            [self.outsider.email, "SIS7002", "AY", "JUL_DEC", "Coordinator", "", "Order 2"],
        ])
        self.assertEqual([r["status"] for r in b["report"]], ["VALID", "INVALID"])
        self.commit(self.acad, b["id"], content, skip=True)
        self.assertTrue(FacultySubjectAssignment.objects.filter(faculty=fac, section=self.sec_a,
                                                                basis="Order 1").exists())


# --- authorization, provisioning, templates, links ------------------------------------------------------------------

class ImportSecurityTests(ImportFixture):
    FACULTY = [["Name", "Email", "Designation", "Department"], ["Dr Sec", "sec@dypiu.ac.in", "Professor", "D1"]]

    def test_unauthorized_imports_are_denied_and_audited(self):
        content = csv_bytes(self.FACULTY)
        for person in (self.phd, self.coe, self.coord, self.p_sch1, self.dept_admin):
            self.assertEqual(self.analyze(person, "faculty", content).status_code, 403)
            self.assertEqual(self.preview(person, "faculty", content, mapping_={"0": "name", "1": "email",
                                                                                "2": "designation",
                                                                                "3": "department"}).status_code, 403)
        ev = AuditEvent.objects.latest("id")
        self.assertEqual((ev.action, ev.allowed), ("imports.faculty", False))
        self.assertFalse(Faculty.objects.filter(email="sec@dypiu.ac.in").exists())
        self.assertEqual(self.client_for(self.p_sch1).get(API + "types/").status_code, 403)

    def test_provisioning_requires_person_management(self):
        content = csv_bytes([SCHOLAR_HEADER, scholar_row("PROV01")])
        r = self.preview(self.phd, "scholars", content, provision="login")   # PhD Cell cannot manage persons
        self.assertEqual(r.status_code, 403)
        r = self.preview(self.acad, "courses", csv_bytes([["code", "title", "credits", "category"],
                                                          ["PC1", "T", "1", "ELECTIVE"]]), provision="login")
        self.assertEqual(r.status_code, 400)  # courses are not people

    def test_row_level_authorization_still_applies(self):
        """A department admin passes the batch gate for assignments but each row is authorized on its offering."""
        fac = self.outsider.faculty_profile
        fac.email = self.outsider.email
        fac.save()
        content = csv_bytes([["Faculty e-mail", "Course code", "Academic year", "Term", "Role", "Section", "Basis"],
                             [self.outsider.email, "SIS7002", "AY", "JUL_DEC", "INSTRUCTOR", "A", "Order 1"]])
        b = self.preview(self.dept_admin, "assignments", content, mapping_={
            "0": "faculty_email", "1": "course_code", "2": "academic_year", "3": "term", "4": "role", "5": "section",
            "6": "basis"}).json()
        self.assertEqual(b["report"][0]["status"], "INVALID")
        self.assertIn("Not authorized", b["report"][0]["reasons"][0])
        self.assertTrue(AuditEvent.objects.filter(action="imports.assignments.row_denied", allowed=False).exists())

    def test_spreadsheet_cannot_grant_capabilities(self):
        content = csv_bytes([["Name", "Email", "Designation", "Department", "Capability", "Role", "Workspace"],
                             ["Dr Grab", "grab@dypiu.ac.in", "Professor", "D1", "SYSTEM_ADMIN", "ACADEMIC_ADMIN",
                              "coe"]])
        a = self.analyze(self.acad, "faculty", content).json()
        unmapped = {c["header"] for c in a["columns"] if not c["field"]}
        self.assertEqual(unmapped, {"Capability", "Role", "Workspace"})
        grants = CapabilityAssignment.objects.count()
        b = self.preview(self.acad, "faculty", content, provision="login").json()
        self.assertEqual(self.commit(self.acad, b["id"], content).status_code, 200)
        self.assertEqual(CapabilityAssignment.objects.count(), grants)
        self.assertEqual([w["key"] for w in workspaces_of(Person.objects.get(email="grab@dypiu.ac.in"))], ["faculty"])

    def test_provisioning_never_merges_existing_people(self):
        Person.objects.create(full_name="Existing", email="clash@dypiu.ac.in")
        content = csv_bytes([self.FACULTY[0], ["Dr Clash", "clash@dypiu.ac.in", "Professor", "D1"],
                             ["Dr Fine", "fine@dypiu.ac.in", "Professor", "D1"]])
        b = self.preview(self.acad, "faculty", content, provision="person").json()
        self.assertEqual([r["status"] for r in b["report"]], ["INVALID", "VALID"])
        self.assertIn("already exists", b["report"][0]["reasons"][0])
        self.commit(self.acad, b["id"], content, skip=True)
        self.assertFalse(Faculty.objects.filter(email="clash@dypiu.ac.in").exists())  # the whole row rolled back
        self.assertIsNone(Person.objects.get(email="clash@dypiu.ac.in").faculty_profile_id)

    def test_activation_links_csv(self):
        content = csv_bytes([self.FACULTY[0], ["=HYPERLINK(\"http://evil\")", "evil@dypiu.ac.in", "Professor", "D1"],
                             ["Dr Link", "link@dypiu.ac.in", "Professor", "D1"]])
        b = self.preview(self.acad, "faculty", content, provision="login").json()
        self.commit(self.acad, b["id"], content)
        self.assertEqual(self.client_for(self.phd).post(f"{API}{b['id']}/activation-links/").status_code, 403)
        r = self.client_for(self.acad).post(f"{API}{b['id']}/activation-links/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Cache-Control"], "no-store")
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["name", "email", "activation_link"])
        self.assertEqual(len(rows), 3)
        self.assertTrue(rows[1][0].startswith("'="))              # formula injection neutralised
        link = rows[2][2]
        uid, token = link.split("uid=")[1].split("&token=")
        ok = APIClient().post("/api/auth/activate/", {"uid": uid, "token": token, "password": "Link-Pass-2026"},
                              format="json")
        self.assertEqual(ok.status_code, 200)
        self.assertNotIn(token, repr(list(AuditEvent.objects.values("after", "reasons"))))

    def test_templates(self):
        import openpyxl
        r = self.client_for(self.phd).get(API + "templates/scholars/")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual([c.value for c in wb["Data"][1]], [f.name for f in TYPES["scholars"].fields])
        self.assertEqual(wb["Data"].max_row, 1)                    # no example rows to import by accident
        fields = {row[0]: row for row in wb["Instructions"].iter_rows(min_row=6, values_only=True)}
        self.assertEqual(fields["prn"][2], "Required")
        self.assertIn("FT (Full-time)", fields["category"][5])
        csv_r = self.client_for(self.phd).get(API + "templates/scholars/", {"kind": "csv"})
        self.assertEqual(csv_r.content.decode("utf-8-sig").strip(), ",".join(f.name for f in TYPES["scholars"].fields))
        self.assertEqual(self.client_for(self.phd).get(API + "templates/faculty/").status_code, 403)

    def test_history_visibility(self):
        self.preview(self.acad, "faculty", csv_bytes(self.FACULTY))
        self.preview(self.phd, "scholars", csv_bytes([SCHOLAR_HEADER, scholar_row("HIST1")]))
        acad = {b["import_type"] for b in self.client_for(self.acad).get(API).json()["results"]}
        phd = {b["import_type"] for b in self.client_for(self.phd).get(API).json()["results"]}
        self.assertEqual((acad, phd), ({"faculty", "scholars"}, {"scholars"}))
        fac_batch = ImportBatch.objects.get(import_type="faculty")
        self.assertEqual(self.client_for(self.phd).get(f"{API}{fac_batch.pk}/").status_code, 403)
        types = {t["key"] for t in self.client_for(self.phd).get(API + "types/").json()["results"]}
        self.assertEqual(types, {"scholars", "enrollments"})
