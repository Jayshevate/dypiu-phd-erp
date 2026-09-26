"""Downloadable import templates (Step A3). The data sheet holds only the header
row (so no example can be imported by accident); the Instructions sheet documents
every field: required or optional, format, example and allowed values."""
import csv
import io

from .types import ImportType


def csv_template(itype: ImportType) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerow([f.name for f in itype.fields])
    return buf.getvalue().encode("utf-8-sig")


def xlsx_template(itype: ImportType) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    data.append([f.name for f in itype.fields])
    header_fill = PatternFill("solid", fgColor="0B1B3D")
    for cell in data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        data.column_dimensions[cell.column_letter].width = max(14, len(str(cell.value)) + 4)
    data.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    info.append([f"{itype.label} import"])
    info["A1"].font = Font(bold=True, size=13)
    info.append([itype.description])
    info.append(["Fill in the Data sheet: one row per record, keep the header row. Existing records are never "
                 "changed; a preview shows every row's outcome before anything is written."])
    info.append([])
    info.append(["Field (column header)", "Label", "Required", "Format", "Example", "Allowed values", "Notes"])
    for cell in info[5]:
        cell.font = Font(bold=True)
    for f in itype.fields:
        info.append([f.name, f.label, "Required" if f.required else "Optional", f.format_text(), f.example,
                     "; ".join(f.allowed()), f.help])
    for col, width in zip("ABCDEFG", (24, 26, 10, 44, 26, 60, 44)):
        info.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
